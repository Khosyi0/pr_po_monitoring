"""
etl_kondisi_stock_bb.py - ETL: Sync data Excel Balans Pupuk & Bahan Baku,
serta Rencana Kebutuhan Bahan Baku, ke PostgreSQL.

Modul ini menangani DUA sumber data yang SELALU diupload sebagai file Excel
TERPISAH (bukan 1 file dengan banyak sheet), tapi diproses lewat SATU kali
pemanggilan run_etl() supaya UI di halaman Manajemen Data cukup punya 1
tombol "Jalankan ETL":

  1. File UTAMA (Config.EXCEL_FILE) -- berisi:
       - Sheet 'Pupuk'       (balans produksi/pengadaan/penjualan pupuk jadi)
       - Sheet 'Bahan Baku'  (balans produksi/pengadaan/pemakaian bahan baku)
       - Sheet 'Stock Chart' (data harian Safety Stock vs Stock PG -- OPSIONAL)

  2. File RENCANA BB (Config.RENCANA_BB_FILE) -- berisi:
       - Sheet 'Data BB <tahun>' (rencana kedatangan/kebutuhan shipment bahan
         baku per komoditas/origin/suplier/bulan) -- OPSIONAL. Nama sheet
         dibentuk otomatis dari Config.TAHUN_DATA (mis. tahun 2026 ->
         sheet 'Data BB 2026'), TIDAK di-hardcode terpisah, karena pasti
         berganti setiap tahun mengikuti judul sheet sumber.

Parsing 'Stock Chart' dan 'Data BB <tahun>' bersifat OPSIONAL: kalau file
kedua tidak diupload, sheet tidak ditemukan, atau gagal diparse, proses ETL
utama (Pupuk + Bahan Baku) TETAP dianggap berhasil -- errornya dilaporkan di
log tapi tidak menghentikan proses, sama seperti pola Stock Chart yang sudah
ada sebelumnya.

Cara pakai:
  1. Sesuaikan Config.EXCEL_FILE, Config.RENCANA_BB_FILE (boleh None kalau
     file kedua tidak diupload), dan Config.TAHUN_DATA (WAJIB diisi manual --
     tidak dideteksi otomatis dari judul sheet, karena format judul bisa
     berubah sewaktu-waktu)
  2. Jalankan: python etl_kondisi_stock_bb.py

Format tabel: LONG (bulan & tahun sebagai kolom data, bukan nama kolom).
  Alasan: rekapan tahun berjalan bisa mulai dari bulan berapa saja (mis.
  2026 mulai April, tahun depan bisa mulai Januari) -- kalau bulan dijadikan
  nama kolom (nilai_apr, nilai_mei, dst) maka skema tabel harus diubah setiap
  pola bulan berbeda. Dengan format long, kolom 'bulan' dan 'tahun_data'
  cukup jadi DATA, sehingga skema tabel tidak perlu berubah kapan pun.

Cara kerja (DELETE per tahun + REINSERT), berlaku utk SEMUA tabel di bawah:
  - Setiap kali ETL dijalankan, HANYA baris dengan tahun_data yang SAMA
    dengan Config.TAHUN_DATA yang dihapus dari tabel terkait, lalu diisi
    ulang dari file yang diupload. Data tahun-tahun lain (histori) TIDAK
    ikut terhapus dan tetap aman di database -- mirip strategi
    sync_sips_data() di etl_sips.py yang DELETE per (bulan_import,
    tahun_import).
  - Kolom 'kategori' pada kondisi_stock_bb_raw TIDAK diisi oleh ETL ini.
    Kolom tersebut dikelola manual lewat UI (v_kondisi_stock_bb.py) untuk
    menandai baris mana yang termasuk kategori standar (Stock Awal,
    Produksi, Pemakaian, Stok Akhir, dll) -- karena label baris di sumber
    sangat bervariasi antar produk dan tidak bisa diklasifikasi otomatis
    secara andal.

===========================================================================
BAGIAN 1: PARSING SHEET 'Pupuk' / 'Bahan Baku' (format balans rapi)
===========================================================================

Cara kerja parsing (per sheet):
  1. Cari baris header (sel kolom A bernilai persis 'No'). Baris tepat di
     bawahnya adalah baris nama bulan (APR, MEI, ..., DES).
  2. Baris "header produk baru" dideteksi jika: kolom A berisi bilangan BULAT
     (format '1.', '12.0', '12', dst -- BUKAN desimal pecahan seperti '0.02')
     DAN kolom B (nama produk) terisi teks, DAN angkanya >= nomor urut produk
     terakhir yang sudah diterima (menoleransi data sumber yang kadang salah
     penomoran/duplikat, seperti dua produk yang sama-sama bernomor '23.').
     Baris seperti 'Buffer Stock ...' yang nyasar ke kolom A tapi kolom B
     kosong TIDAK dianggap header produk baru.
  3. Semua baris di antara satu header produk dan header produk berikutnya
     (yang kolom B/label-nya terisi) disimpan apa adanya sebagai baris data
     milik produk tsb, termasuk baris breakdown berprefix '►'.
  4. Kolom C (Realisasi) disimpan sebagai baris dengan bulan='realisasi',
     urutan_bulan=0. Kolom D-L (Prognosa per bulan) masing-masing jadi satu
     baris dengan bulan sesuai nama bulannya (apr, mei, ..., des) dan
     urutan_bulan 1-12 mengikuti urutan KALENDER (bukan urutan kolom di
     Excel), agar render selalu bisa diurutkan Jan->Des apa pun urutannya.
  5. Nilai '-' di Excel disimpan sebagai NULL (bukan 0), karena '-' berarti
     "tidak ada transaksi", beda makna dengan nilai nol.

===========================================================================
BAGIAN 2: PARSING SHEET 'Stock Chart' (data harian, OPSIONAL)
===========================================================================

Sheet ini beda total polanya dari 'Pupuk'/'Bahan Baku': setiap blok virtual
produk punya kolom Tanggal (datetime asli) + Safety Stock + Stock PG
berdampingan secara horizontal. Lihat CHART_MAPPING di bawah untuk detail
kolom per virtual produk.

===========================================================================
BAGIAN 3: PARSING SHEET 'Data BB <tahun>' (rencana shipment, OPSIONAL)
===========================================================================

Sheet ini formatnya jauh lebih bebas dibanding Bagian 1:
  - Setiap komoditas adalah satu BLOK baris (mis. 'Phosphate Rock' baris
    4-17, 'Sulphur' baris 18-28, dst). Blok terdeteksi dari kolom B
    (Komoditas) terisi teks -- baik hasil merged cell (komoditas dengan
    banyak baris origin/suplier) maupun baris tunggal.
  - Tiap baris dalam blok = kombinasi Origin + Nama Suplier + Harga. Satu
    komoditas bisa punya banyak baris origin/suplier berbeda.
  - Kolom F..Q (12 bulan Jan..Des) isinya TEKS BEBAS multi-baris berisi info
    shipment (tanggal, nama kapal, qty, kadang PO/Laycan/ETA/Depart). Tidak
    ada format baku, jadi TIDAK di-parse jadi field terstruktur -- disimpan
    apa adanya sebagai teks.
  - Kolom R..X (Estimasi Qty/shipment, Harga Perolehan Terakhir, Posisi
    Stok, Safety Stock, Keterangan, Rute Kapal Terdampak, Mitigasi Risiko)
    adalah info ringkas milik KOMODITAS (bukan milik baris origin/suplier
    tertentu), biasanya hanya terisi di baris pertama tiap blok (kadang di
    baris kedua utk komoditas dgn banyak sub-harga, mis. Sulphuric Acid).
    Nilai ini di-"forward fill" ke seluruh baris dalam blok yang sama.
  - Baris berisi harga tapi TIDAK ada Origin/Suplier maupun shipment apapun
    tetap disimpan (bukan dilewati) supaya tampilan data mentah tetap utuh.
  - Baris "Keterangan"/legenda warna di akhir sheet BUKAN data komoditas dan
    tidak diproses; parsing berhenti begitu ada baris dengan kolom B ==
    'Keterangan'.

Requirements:
  pip install pandas openpyxl sqlalchemy psycopg2-binary --break-system-packages
"""

import re
import pandas as pd
import numpy as np
from datetime import datetime, date
from sqlalchemy import create_engine, text
import warnings
import os
warnings.filterwarnings('ignore')


# =====================================================================
# KONFIGURASI
# =====================================================================

class Config:
    DB_HOST     = 'localhost'
    DB_PORT     = '5432'
    DB_NAME     = 'pr_po_monitoring'
    DB_USER     = 'postgres'
    DB_PASSWORD = 'Hx4Khos2'

    # -- File UTAMA: sheet 'Pupuk', 'Bahan Baku', 'Stock Chart' --
    EXCEL_FILE   = 'Data_Outline_Laporan.xlsx'
    SHEET_PUPUK  = 'Pupuk'
    SHEET_BB     = 'Bahan Baku'
    SHEET_CHART  = 'Stock Chart'

    # -- File RENCANA BB (terpisah, OPSIONAL): sheet 'Data BB <tahun>' --
    # Diisi None kalau file kedua ini tidak diupload -- ETL utama tetap
    # berjalan normal tanpa memproses bagian ini.
    RENCANA_BB_FILE = None

    # WAJIB diisi manual sebelum run_etl() dipanggil -- tahun data yang
    # sedang direkap di file ini (mis. 2026). Dipilih user lewat dropdown
    # di halaman Manajemen Data, bukan dideteksi otomatis dari judul sheet.
    # Dipakai juga untuk membentuk nama sheet 'Data BB <tahun>' pada file
    # RENCANA_BB_FILE.
    TAHUN_DATA   = None


# =====================================================================
# DATABASE
# =====================================================================

def db_get_engine():
    cs = (f"postgresql://{Config.DB_USER}:{Config.DB_PASSWORD}"
          f"@{Config.DB_HOST}:{Config.DB_PORT}/{Config.DB_NAME}")
    return create_engine(cs)


# =====================================================================
# HELPERS UMUM
# =====================================================================

# Bilangan BULAT saja (boleh format '1.', '12.0', '12') -- BUKAN desimal
# pecahan seperti '0.02' yang muncul di baris non-header (mis. "Pemakaian
# sbg Filler NPK (2%)").
PRODUK_PATTERN = re.compile(r'^(\d+)(\.0)?\.?$')

# Mapping nama bulan (versi disingkat/dinormalisasi dari header Excel) ->
# urutan kalender 1-12. Dipakai supaya render selalu bisa diurutkan
# Jan->Des apa pun urutan kolom aslinya di sheet.
BULAN_URUTAN = {
    'jan': 1, 'feb': 2, 'mar': 3, 'apr': 4, 'mei': 5, 'jun': 6,
    'jul': 7, 'agust': 8, 'ags': 8, 'agu': 8, 'sep': 9,
    'okt': 10, 'nop': 11, 'nov': 11, 'des': 12,
}

# Urutan bulan kalender dengan key yang dipakai di tabel rencana_bb_raw
# (sengaja daftar terpisah dari BULAN_URUTAN karena sheet Data BB tidak
# perlu menormalisasi nama bulan dari header -- kolomnya sudah tetap F..Q).
BULAN_URUTAN_RENCANA_BB = [
    ('jan', 1), ('feb', 2), ('mar', 3), ('apr', 4), ('mei', 5), ('jun', 6),
    ('jul', 7), ('agust', 8), ('sep', 9), ('okt', 10), ('nop', 11), ('des', 12),
]


def clean_nilai(v):
    """Konversi nilai sel Excel -> float atau None. '-' dan sejenisnya -> None."""
    if v is None:
        return None
    if isinstance(v, (int, float)):
        try:
            if pd.isna(v):
                return None
        except Exception:
            pass
        return float(v)
    s = str(v).strip()
    if s == '' or s == '-':
        return None
    s_clean = s.replace('.', '').replace(',', '.')
    try:
        return float(s_clean)
    except Exception:
        return None


def clean_text(v):
    """Konversi nilai sel -> string bersih, atau None kalau kosong.
    Nilai datetime disimpan sbg string tanggal (DD-MM-YYYY) supaya konsisten
    dgn sel lain yang isinya teks tanggal manual. Dipakai khusus utk parsing
    sheet 'Data BB <tahun>' yang selnya campur teks/angka/tanggal."""
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return v.strftime('%d-%m-%Y')
    if isinstance(v, float):
        if v == int(v):
            return str(int(v))
        return str(v)
    s = str(v).strip()
    return s if s != '' else None


# =====================================================================
# BAGIAN 1: PARSING SHEET 'Pupuk' / 'Bahan Baku'
# =====================================================================

def parse_sheet(ws, sumber_sheet, tahun_data):
    """
    Parse satu worksheet (openpyxl) 'Pupuk' atau 'Bahan Baku' menjadi list of
    dict baris data mentah FORMAT LONG (1 baris label x 1 bulan = 1 record),
    mengikuti struktur balans produksi/pengadaan/penjualan yang dijelaskan
    di docstring modul ini.
    """
    # 1. Cari baris header ('No' persis di kolom A)
    header_row = None
    for r in range(1, 15):
        if ws.cell(row=r, column=1).value == 'No':
            header_row = r
            break
    if header_row is None:
        raise ValueError(
            f"Baris header 'No' tidak ditemukan di sheet '{sumber_sheet}' "
            f"(dicek baris 1-14 kolom A). Struktur sheet mungkin berubah."
        )

    month_row = header_row + 1

    # 2. Kolom bulan PROGNOSA: D..L (index 4-12). Kolom C (Realisasi)
    #    SENGAJA ditangani terpisah -- walau nama bulannya bisa kebetulan
    #    sama dengan salah satu kolom prognosa (mis. 'Realisasi s/d JUN' vs
    #    kolom Prognosa 'JUN') -- supaya tidak tertukar/duplikat makna.
    col_bulan_info = {}  # {col_index: (bulan_key, urutan_bulan_kalender)}
    for c in range(4, 13):
        v = ws.cell(row=month_row, column=c).value
        if v:
            name_raw = str(v).strip().lower()
            key = re.sub(r'[^a-z]', '', name_raw)
            urutan = BULAN_URUTAN.get(key)
            if urutan is None:
                raise ValueError(
                    f"Nama bulan '{v}' di kolom {c} sheet '{sumber_sheet}' "
                    f"tidak dikenali. Cek BULAN_URUTAN di skrip ETL."
                )
            col_bulan_info[c] = (key, urutan)
    if not col_bulan_info:
        raise ValueError(
            f"Kolom bulan tidak ditemukan di baris {month_row} sheet "
            f"'{sumber_sheet}'. Struktur sheet mungkin berubah."
        )

    col_realisasi = 3   # C
    col_total     = 13  # M -- 'PROGNOSA JAN s/d DES'
    col_rkap      = 14  # N -- 'ASUMSI RKAP AWAL <tahun>'

    # 3. Iterasi baris data
    records = []
    current_produk = None
    current_urutan_produk = None
    urutan_baris_in_produk = 0
    last_num = 0

    max_row = ws.max_row
    for r in range(month_row + 1, max_row + 1):
        a_val = ws.cell(row=r, column=1).value
        b_val = ws.cell(row=r, column=2).value

        is_new_produk = False
        angka = None
        if a_val is not None and b_val is not None and str(b_val).strip() != '':
            a_str = str(a_val).strip()
            m = PRODUK_PATTERN.match(a_str)
            if m:
                angka = int(m.group(1))
                # terima jika nomor naik (wajar) atau >= terakhir (toleransi
                # duplikat/typo penomoran di sumber, mis. dua produk sama2 '23.')
                if angka >= last_num:
                    is_new_produk = True

        if is_new_produk:
            current_produk = str(b_val).strip()
            current_urutan_produk = angka
            last_num = angka
            urutan_baris_in_produk = 0
            continue  # baris header produk sendiri tidak disimpan sbg baris data

        if current_produk is None:
            continue  # belum masuk produk manapun (baris sebelum produk pertama)
        if b_val is None or str(b_val).strip() == '':
            continue  # baris kosong / separator, dilewati

        label = str(b_val).strip()
        urutan_baris_in_produk += 1

        nilai_total = clean_nilai(ws.cell(row=r, column=col_total).value)
        nilai_rkap  = clean_nilai(ws.cell(row=r, column=col_rkap).value)

        base = dict(
            sumber_sheet=sumber_sheet,
            produk=current_produk,
            urutan_produk=current_urutan_produk,
            urutan_baris=urutan_baris_in_produk,
            label_baris=label,
            tahun_data=tahun_data,
            nilai_total_jan_des=nilai_total,
            nilai_rkap=nilai_rkap,
        )

        # baris 'realisasi' (urutan_bulan = 0, selalu paling awal saat di-sort)
        records.append({
            **base,
            'bulan': 'realisasi',
            'urutan_bulan': 0,
            'nilai': clean_nilai(ws.cell(row=r, column=col_realisasi).value),
        })

        # baris tiap bulan prognosa
        for c, (bkey, urut) in col_bulan_info.items():
            records.append({
                **base,
                'bulan': bkey,
                'urutan_bulan': urut,
                'nilai': clean_nilai(ws.cell(row=r, column=c).value),
            })

    return records


def load_and_parse():
    if not os.path.exists(Config.EXCEL_FILE):
        raise FileNotFoundError(f"File tidak ditemukan: '{Config.EXCEL_FILE}'")
    if not Config.TAHUN_DATA:
        raise ValueError(
            "Config.TAHUN_DATA belum diisi. Tahun data harus dipilih manual "
            "sebelum menjalankan ETL ini (mis. lewat dropdown di halaman "
            "Manajemen Data)."
        )

    import openpyxl
    wb = openpyxl.load_workbook(Config.EXCEL_FILE, data_only=True)

    all_records = []
    for sheet_name, label in [(Config.SHEET_PUPUK, 'Pupuk'), (Config.SHEET_BB, 'Bahan Baku')]:
        if sheet_name not in wb.sheetnames:
            raise ValueError(
                f"Sheet '{sheet_name}' tidak ditemukan di file. "
                f"Sheet yang ada: {wb.sheetnames}"
            )
        ws = wb[sheet_name]
        recs = parse_sheet(ws, label, Config.TAHUN_DATA)
        n_baris_label = len(set((r['produk'], r['urutan_baris']) for r in recs))
        print(f"   Sheet '{sheet_name}' : {n_baris_label:,} baris label -> {len(recs):,} record")
        all_records.extend(recs)

    if not all_records:
        print("   ⚠️ Peringatan: Tidak ada baris yang berhasil di-parse.")
        return pd.DataFrame()

    df = pd.DataFrame(all_records)
    print(f"   Total gabungan : {len(df):,} record (Pupuk + Bahan Baku, tahun {Config.TAHUN_DATA})")
    return df


# =====================================================================
# BAGIAN 2: CHART HARIAN (sheet 'Stock Chart') -- OPSIONAL
# =====================================================================
# Sheet ini berbeda total polanya dari 'Pupuk'/'Bahan Baku': setiap blok
# virtual produk punya kolom Tanggal (datetime asli) + Safety Stock + Stock
# PG berdampingan secara horizontal. Beberapa blok (KCL) berbagi satu kolom
# Tanggal untuk 2 sub-jenis (Merah & Putih), makanya struktur mapping di
# bawah berbasis nomor kolom eksplisit per virtual produk, bukan dideteksi
# otomatis dari header (karena headernya sendiri sudah tidak konsisten,
# mis. sebagian sub-blok tidak punya kolom 'Tanggal' sendiri).
#
# Kolom dihitung 1-based (A=1, B=2, ...), sesuai openpyxl.

CHART_MAPPING = {
    "Phospate Rock": {"tanggal": 9,  "items": [{"jenis": "default", "safety": 10, "stock": 11}]},
    "ZA":            {"tanggal": 13, "items": [{"jenis": "default", "safety": 14, "stock": 15}]},
    "Asam Sulfat":   {"tanggal": 17, "items": [{"jenis": "default", "safety": 18, "stock": 19}]},
    "Sulphur":       {"tanggal": 21, "items": [{"jenis": "default", "safety": 22, "stock": 23}]},
    "KCL":           {"tanggal": 25, "items": [
        {"jenis": "Merah", "safety": 26, "stock": 27},
        {"jenis": "Putih", "safety": 29, "stock": 30},
    ]},
    "DAP":           {"tanggal": 32, "items": [{"jenis": "default", "safety": 33, "stock": 34}]},
    "NH4Cl":         {"tanggal": 36, "items": [{"jenis": "default", "safety": 37, "stock": 38}]},
    "Ammonia":       {"tanggal": 40, "items": [{"jenis": "default", "safety": 41, "stock": 42}]},
    "PA":            {"tanggal": 44, "items": [{"jenis": "default", "safety": 45, "stock": 46}]},
}

CHART_HEADER_ROW = 2  # baris sub-header ('Tanggal', 'Safety Stock', dst); data mulai baris berikutnya


def clean_tanggal(v):
    """Konversi sel Excel (biasanya sudah datetime asli) -> date, atau None."""
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except Exception:
        pass
    if hasattr(v, 'date'):
        return v.date()
    try:
        return pd.to_datetime(v, dayfirst=True).date()
    except Exception:
        return None


def parse_chart_harian(ws, tahun_data):
    """Parse sheet 'Stock Chart' -> list of dict, 1 baris per (virtual_produk, jenis, tanggal)."""
    records = []
    for virtual_produk, cfg in CHART_MAPPING.items():
        col_tanggal = cfg["tanggal"]
        for r in range(CHART_HEADER_ROW + 1, ws.max_row + 1):
            tgl_val = ws.cell(row=r, column=col_tanggal).value
            tgl = clean_tanggal(tgl_val)
            if tgl is None:
                continue
            for item in cfg["items"]:
                safety = clean_nilai(ws.cell(row=r, column=item["safety"]).value)
                stock = clean_nilai(ws.cell(row=r, column=item["stock"]).value)
                records.append({
                    "virtual_produk": virtual_produk,
                    "jenis": item["jenis"],
                    "tanggal": tgl,
                    "safety_stock": safety,
                    "stock_pg": stock,
                    "tahun_data": tahun_data,
                })
    return records


def load_and_parse_chart():
    """Parse sheet Config.SHEET_CHART. Return DataFrame (bisa kosong jika sheet tidak ada)."""
    import openpyxl
    wb = openpyxl.load_workbook(Config.EXCEL_FILE, data_only=True)

    if Config.SHEET_CHART not in wb.sheetnames:
        print(f"   ⚠️ Sheet '{Config.SHEET_CHART}' tidak ditemukan -- lewati chart harian.")
        return pd.DataFrame()

    ws = wb[Config.SHEET_CHART]
    records = parse_chart_harian(ws, Config.TAHUN_DATA)

    if not records:
        print(f"   ⚠️ Tidak ada data harian ter-parse dari sheet '{Config.SHEET_CHART}'.")
        return pd.DataFrame()

    df = pd.DataFrame(records)
    n_produk = df['virtual_produk'].nunique()
    n_tanggal = df['tanggal'].nunique()
    print(f"   Sheet '{Config.SHEET_CHART}' : {n_produk} virtual produk x ~{n_tanggal} tanggal -> {len(df):,} record")
    return df


# =====================================================================
# BAGIAN 3: RENCANA KEBUTUHAN BB (sheet 'Data BB <tahun>', file terpisah, OPSIONAL)
# =====================================================================

# Kolom bulan F..Q (index 1-based openpyxl): F=6 ... Q=17
RBB_KOLOM_BULAN_MULAI = 6

# Kolom lain (1-based) khusus sheet 'Data BB <tahun>'
RBB_KOL_KOMODITAS = 2   # B
RBB_KOL_ORIGIN    = 3   # C
RBB_KOL_SUPLIER   = 4   # D
RBB_KOL_HARGA     = 5   # E
RBB_KOL_QTY_SHP   = 18  # R  - Estimasi Qty/shipment
RBB_KOL_HARGA_AKH = 19  # S  - Harga Perolehan Terakhir
RBB_KOL_POSISI    = 20  # T  - Posisi Stok
RBB_KOL_SAFETY    = 21  # U  - Safety Stock
RBB_KOL_KET       = 22  # V  - Keterangan
RBB_KOL_RUTE      = 23  # W  - Rute Kapal Terdampak
RBB_KOL_MITIGASI  = 24  # X  - Mitigasi Risiko

RBB_DATA_START_ROW = 4  # baris pertama data komoditas di sheet 'Data BB <tahun>'


def _rbb_is_stop_row(komoditas_val):
    """Baris 'Keterangan' dan sesudahnya menandai akhir data komoditas
    (legenda warna font, bukan data)."""
    if komoditas_val is None:
        return False
    return str(komoditas_val).strip().lower() == 'keterangan'


def _rbb_build_merge_lookup(ws, kolom_list):
    """Bangun lookup {(row, col): nilai_asli_merge} untuk kolom2 tertentu.

    PENTING: Kolom Komoditas (B), Origin (C), Suplier (D), dan Harga (E) di
    sheet 'Data BB <tahun>' SEMUANYA bisa di-merge secara vertikal, tapi
    dengan POLA MERGE YANG BERBEDA-BEDA per kolom (mis. Origin/Suplier
    di-merge per 2 baris, sedangkan Harga di-merge per 3-4 baris sesuai
    kelompok harga yang sama -- lihat contoh 'Phosphate Rock': origin Egypt
    MG punya 3 suplier berbeda [Agrifields/Upstream/Midgulf] yang SEMUANYA
    berbagi 1 sel harga 112.4 ter-merge, padahal Origin & Suplier sendiri
    tidak ikut ter-merge sepanjang itu).

    openpyxl hanya mengisi value pada sel PALING KIRI-ATAS dari suatu merged
    range; sel-sel lain dalam range yang sama bernilai None walau
    Excel menampilkannya seolah terisi. Kalau dibaca apa adanya per sel
    tanpa memperhitungkan ini, baris ke-2/3/dst dalam 1 merge range akan
    tampak "kosong" padahal sebenarnya berbagi nilai yang sama.

    Fungsi ini membangun lookup eksplisit dari seluruh merged range pada
    kolom-kolom yang diminta, supaya setiap sel (termasuk yang secara
    teknis None karena bagian dari merge) bisa di-resolve ke nilai yang
    benar SEBELUM proses parsing baris dimulai.
    """
    lookup = {}
    for mc in ws.merged_cells.ranges:
        if mc.min_col in kolom_list and mc.min_col == mc.max_col:
            nilai = ws.cell(row=mc.min_row, column=mc.min_col).value
            for r in range(mc.min_row, mc.max_row + 1):
                lookup[(r, mc.min_col)] = nilai
    return lookup


def _rbb_get_cell_value(ws, merge_lookup, row, col):
    """Ambil value sel, tapi utamakan hasil resolve dari merge_lookup kalau
    sel ini bagian dari suatu merged range (lihat _rbb_build_merge_lookup)."""
    if (row, col) in merge_lookup:
        return merge_lookup[(row, col)]
    return ws.cell(row=row, column=col).value


def parse_sheet_rencana_bb(ws, tahun_data):
    """
    Parse sheet 'Data BB <tahun>' -> list of dict FORMAT LONG:
    1 baris = 1 kombinasi (komoditas, origin/suplier/harga pada urutan_baris
    tertentu) x 1 bulan.
    """
    max_row = ws.max_row

    # Kolom Komoditas/Origin/Suplier/Harga bisa di-merge vertikal dengan pola
    # berbeda2 -- resolve semua merge-nya lebih dulu (lihat docstring
    # _rbb_build_merge_lookup di atas).
    merge_lookup = _rbb_build_merge_lookup(
        ws, [RBB_KOL_KOMODITAS, RBB_KOL_ORIGIN, RBB_KOL_SUPLIER, RBB_KOL_HARGA]
    )

    records = []
    current_komoditas = None
    current_urutan_komoditas = 0
    urutan_baris_in_komoditas = 0

    # info ringkas (kolom R..X) di-forward-fill dalam 1 blok komoditas
    ringkas = {
        'estimasi_qty_shipment': None,
        'harga_perolehan_terakhir': None,
        'posisi_stok': None,
        'safety_stock': None,
        'keterangan': None,
        'rute_kapal_terdampak': None,
        'mitigasi_risiko': None,
    }

    for r in range(RBB_DATA_START_ROW, max_row + 1):
        # Komoditas dibaca LANGSUNG dari sel asli (bukan lewat merge_lookup)
        # supaya deteksi "mulai blok baru" tetap akurat berdasarkan baris
        # pertama merge komoditas, bukan hasil resolve yang sudah di-expand.
        komoditas_val = ws.cell(row=r, column=RBB_KOL_KOMODITAS).value

        if _rbb_is_stop_row(komoditas_val):
            break

        komoditas_bersih = clean_text(komoditas_val)
        if komoditas_bersih:
            # mulai blok komoditas baru
            current_komoditas = komoditas_bersih
            current_urutan_komoditas += 1
            urutan_baris_in_komoditas = 0
            ringkas = {k: None for k in ringkas}  # reset forward-fill utk blok baru

        if current_komoditas is None:
            continue  # belum masuk komoditas manapun (baris sebelum data pertama)

        # Origin/Suplier/Harga di-resolve lewat merge_lookup supaya baris ke-2/3/dst
        # dalam 1 merged range (mis. 3 suplier berbeda yg berbagi 1 sel harga)
        # tetap mendapat nilai yang benar, bukan kosong.
        origin  = clean_text(_rbb_get_cell_value(ws, merge_lookup, r, RBB_KOL_ORIGIN))
        suplier = clean_text(_rbb_get_cell_value(ws, merge_lookup, r, RBB_KOL_SUPLIER))
        harga   = clean_text(_rbb_get_cell_value(ws, merge_lookup, r, RBB_KOL_HARGA))

        # ambil nilai ringkas kolom R..X baris ini kalau ada, kalau tidak
        # pertahankan nilai forward-fill dari baris sebelumnya di blok yg sama
        for col_idx, key in [
            (RBB_KOL_QTY_SHP, 'estimasi_qty_shipment'),
            (RBB_KOL_HARGA_AKH, 'harga_perolehan_terakhir'),
            (RBB_KOL_POSISI, 'posisi_stok'),
            (RBB_KOL_SAFETY, 'safety_stock'),
            (RBB_KOL_KET, 'keterangan'),
            (RBB_KOL_RUTE, 'rute_kapal_terdampak'),
            (RBB_KOL_MITIGASI, 'mitigasi_risiko'),
        ]:
            v = clean_text(ws.cell(row=r, column=col_idx).value)
            if v is not None:
                ringkas[key] = v

        bulan_values = {}
        ada_isi_bulan = False
        for i, (bkey, urut) in enumerate(BULAN_URUTAN_RENCANA_BB):
            col = RBB_KOLOM_BULAN_MULAI + i
            isi = clean_text(ws.cell(row=r, column=col).value)
            bulan_values[bkey] = isi
            if isi is not None:
                ada_isi_bulan = True

        # Lewati baris yang benar2 kosong total (tidak ada origin, suplier,
        # harga, maupun isi bulan apapun) -- baris pemisah/kosong di Excel.
        if not any([origin, suplier, harga, ada_isi_bulan]):
            continue

        urutan_baris_in_komoditas += 1

        base = dict(
            tahun_data=tahun_data,
            komoditas=current_komoditas,
            urutan_komoditas=current_urutan_komoditas,
            urutan_baris=urutan_baris_in_komoditas,
            origin=origin,
            suplier=suplier,
            harga=harga,
            estimasi_qty_shipment=ringkas['estimasi_qty_shipment'],
            harga_perolehan_terakhir=ringkas['harga_perolehan_terakhir'],
            posisi_stok=ringkas['posisi_stok'],
            safety_stock=ringkas['safety_stock'],
            keterangan=ringkas['keterangan'],
            rute_kapal_terdampak=ringkas['rute_kapal_terdampak'],
            mitigasi_risiko=ringkas['mitigasi_risiko'],
        )

        for i, (bkey, urut) in enumerate(BULAN_URUTAN_RENCANA_BB):
            records.append({
                **base,
                'bulan': bkey,
                'urutan_bulan': urut,
                'isi_shipment': bulan_values[bkey],
            })

    return records


def load_and_parse_rencana_bb():
    """Parse sheet 'Data BB <tahun>' dari Config.RENCANA_BB_FILE.
    Return DataFrame (bisa kosong jika file/sheet tidak ada atau gagal)."""
    if not Config.RENCANA_BB_FILE:
        print("   ⚠️ Config.RENCANA_BB_FILE tidak diisi -- lewati Rencana Kebutuhan BB.")
        return pd.DataFrame()

    if not os.path.exists(Config.RENCANA_BB_FILE):
        print(f"   ⚠️ File Rencana Kebutuhan BB tidak ditemukan: '{Config.RENCANA_BB_FILE}' -- lewati.")
        return pd.DataFrame()

    sheet_name = f"Data BB {Config.TAHUN_DATA}"

    import openpyxl
    wb = openpyxl.load_workbook(Config.RENCANA_BB_FILE, data_only=True)

    if sheet_name not in wb.sheetnames:
        print(f"   ⚠️ Sheet '{sheet_name}' tidak ditemukan di file Rencana Kebutuhan BB "
              f"(sheet yang ada: {wb.sheetnames}) -- lewati.")
        return pd.DataFrame()

    ws = wb[sheet_name]
    records = parse_sheet_rencana_bb(ws, Config.TAHUN_DATA)

    if not records:
        print(f"   ⚠️ Tidak ada baris yang berhasil di-parse dari sheet '{sheet_name}'.")
        return pd.DataFrame()

    df = pd.DataFrame(records)
    n_komoditas = df['komoditas'].nunique()
    n_baris_label = len(set((r['komoditas'], r['urutan_baris']) for r in records))
    print(f"   Sheet '{sheet_name}' : {n_komoditas} komoditas, "
          f"{n_baris_label:,} baris origin/suplier -> {len(df):,} record")
    return df


# =====================================================================
# LOAD
# =====================================================================

def sync_kondisi_stock_bb(df: pd.DataFrame, engine):
    """DELETE hanya baris tahun_data yang sama, lalu INSERT ulang tahun tsb.
    Data tahun lain di tabel tidak tersentuh/aman."""
    if df.empty:
        return

    tahun = Config.TAHUN_DATA
    with engine.begin() as conn:
        deleted = conn.execute(
            text("DELETE FROM kondisi_stock_bb_raw WHERE tahun_data = :tahun"),
            {'tahun': tahun}
        ).rowcount

    for i in range(0, len(df), 2000):
        chunk = df.iloc[i:i+2000]
        with engine.begin() as conn:
            chunk.to_sql('kondisi_stock_bb_raw', conn, if_exists='append', index=False)

    print(f"   Data          : {deleted} lama dihapus (tahun {tahun} saja) → {len(df):,} baru diinsert (tahun {tahun})")


def sync_chart_harian(df: pd.DataFrame, engine):
    """DELETE hanya baris tahun_data yang sama di kondisi_stock_bb_chart_harian,
    lalu INSERT ulang. Sama pola dengan sync_kondisi_stock_bb, tabel terpisah."""
    if df.empty:
        return

    tahun = Config.TAHUN_DATA
    with engine.begin() as conn:
        deleted = conn.execute(
            text("DELETE FROM kondisi_stock_bb_chart_harian WHERE tahun_data = :tahun"),
            {'tahun': tahun}
        ).rowcount

    for i in range(0, len(df), 2000):
        chunk = df.iloc[i:i+2000]
        with engine.begin() as conn:
            chunk.to_sql('kondisi_stock_bb_chart_harian', conn, if_exists='append', index=False)

    print(f"   Data Chart    : {deleted} lama dihapus (tahun {tahun} saja) → {len(df):,} baru diinsert (tahun {tahun})")


def sync_rencana_bb(df: pd.DataFrame, engine):
    """DELETE hanya baris tahun_data yang sama di rencana_bb_raw, lalu
    INSERT ulang. Sama pola dengan sync_kondisi_stock_bb, tabel terpisah."""
    if df.empty:
        return

    tahun = Config.TAHUN_DATA
    with engine.begin() as conn:
        deleted = conn.execute(
            text("DELETE FROM rencana_bb_raw WHERE tahun_data = :tahun"),
            {'tahun': tahun}
        ).rowcount

    for i in range(0, len(df), 2000):
        chunk = df.iloc[i:i+2000]
        with engine.begin() as conn:
            chunk.to_sql('rencana_bb_raw', conn, if_exists='append', index=False)

    print(f"   Data Rencana BB : {deleted} lama dihapus (tahun {tahun} saja) → {len(df):,} baru diinsert (tahun {tahun})")


# =====================================================================
# MAIN
# =====================================================================

def run_etl():
    print("=" * 55)
    print("🚀 KONDISI STOCK BB - ETL")
    print(f"   File Utama      : {Config.EXCEL_FILE}")
    print(f"   Sheet Utama     : '{Config.SHEET_PUPUK}' + '{Config.SHEET_BB}' + '{Config.SHEET_CHART}'")
    print(f"   File Rencana BB : {Config.RENCANA_BB_FILE or '(tidak diupload)'}")
    print(f"   Tahun Data      : {Config.TAHUN_DATA}")
    print("=" * 55)

    if not Config.TAHUN_DATA:
        print("❌ Config.TAHUN_DATA belum diisi. Batalkan proses ETL.")
        return False

    try:
        engine = db_get_engine()
        with engine.connect() as conn:
            conn.execute(text("SELECT 1"))
        print("✅ Koneksi database OK\n")
    except Exception as e:
        print(f"❌ Koneksi database gagal: {e}")
        return False

    # -- Bagian 1: Pupuk + Bahan Baku (WAJIB, kegagalan disini menggagalkan ETL) --
    try:
        df = load_and_parse()
        print()
        sync_kondisi_stock_bb(df, engine)
    except FileNotFoundError as e:
        print(f"\n❌ {e}")
        return False
    except Exception as e:
        import traceback
        print(f"\n❌ Error: {e}")
        traceback.print_exc()
        return False

    # -- Bagian 2: Chart harian (OPSIONAL) --
    try:
        print()
        df_chart = load_and_parse_chart()
        sync_chart_harian(df_chart, engine)
    except Exception as e:
        import traceback
        print(f"\n⚠️ Gagal memproses data Chart Harian (data Pupuk/Bahan Baku tetap tersimpan): {e}")
        traceback.print_exc()

    # -- Bagian 3: Rencana Kebutuhan BB (OPSIONAL, file terpisah) --
    try:
        print()
        df_rencana = load_and_parse_rencana_bb()
        sync_rencana_bb(df_rencana, engine)
    except Exception as e:
        import traceback
        print(f"\n⚠️ Gagal memproses data Rencana Kebutuhan BB (data Pupuk/Bahan Baku tetap tersimpan): {e}")
        traceback.print_exc()

    with engine.connect() as conn:
        total = conn.execute(text("SELECT COUNT(*) FROM kondisi_stock_bb_raw")).scalar()
        total_tahun_ini = conn.execute(
            text("SELECT COUNT(*) FROM kondisi_stock_bb_raw WHERE tahun_data = :tahun"),
            {'tahun': Config.TAHUN_DATA}
        ).scalar()
        total_produk = conn.execute(text(
            "SELECT COUNT(DISTINCT sumber_sheet || '|' || produk) FROM kondisi_stock_bb_raw WHERE tahun_data = :tahun"
        ), {'tahun': Config.TAHUN_DATA}).scalar()
        semua_tahun = conn.execute(text(
            "SELECT DISTINCT tahun_data FROM kondisi_stock_bb_raw ORDER BY tahun_data"
        )).fetchall()
        daftar_tahun = ', '.join(str(t[0]) for t in semua_tahun)
        total_chart = conn.execute(text(
            "SELECT COUNT(*) FROM kondisi_stock_bb_chart_harian WHERE tahun_data = :tahun"
        ), {'tahun': Config.TAHUN_DATA}).scalar()
        total_rencana = conn.execute(text(
            "SELECT COUNT(*) FROM rencana_bb_raw WHERE tahun_data = :tahun"
        ), {'tahun': Config.TAHUN_DATA}).scalar()

    print("\n" + "=" * 55)
    print("✅ ETL SELESAI")
    print(f"   Record tahun {Config.TAHUN_DATA}  : {total_tahun_ini:,}")
    print(f"   Produk/item tahun ini : {total_produk:,}")
    print(f"   Record Chart Harian tahun ini : {total_chart:,}")
    print(f"   Record Rencana Kebutuhan BB tahun ini : {total_rencana:,}")
    print(f"   Total record semua tahun di DB : {total:,}")
    print(f"   Tahun yang tersimpan di DB     : {daftar_tahun}")
    print("=" * 55)
    return True


if __name__ == "__main__":
    run_etl()